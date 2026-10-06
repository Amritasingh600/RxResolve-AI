"""
Uploaded document handling: validation, text extraction and simple field extraction.

Supported: .pdf (pypdf), .docx (python-docx), .txt.
Field extraction uses regular expressions that look for "Label: value" lines,
which is what most rejection notices look like. It is intentionally simple.
"""
import io
import re
from datetime import datetime
from pathlib import Path

import config


class DocumentError(Exception):
    """Raised for invalid or unreadable uploads. The message is shown to the user."""


def safe_filename(filename: str) -> str:
    """Strip any directory parts and unusual characters from an uploaded file name."""
    name = Path(filename or "").name
    name = re.sub(r"[^A-Za-z0-9._-]", "_", name).strip("._")
    return name[:100] or "upload"


def validate_upload(filename: str, content: bytes) -> tuple[str, str]:
    """Check the extension, size and basic file signature. Returns (safe_name, extension)."""
    name = safe_filename(filename)
    extension = Path(name).suffix.lower()

    if extension not in config.ALLOWED_UPLOAD_EXTENSIONS:
        allowed = ", ".join(sorted(config.ALLOWED_UPLOAD_EXTENSIONS))
        raise DocumentError(f"File type '{extension or 'unknown'}' is not allowed. Allowed types: {allowed}.")
    if len(content) == 0:
        raise DocumentError("The uploaded file is empty.")
    if len(content) > config.MAX_UPLOAD_BYTES:
        raise DocumentError(f"File is too large. Maximum size is {config.MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")

    # Make sure the content really matches the extension.
    if extension == ".pdf" and not content.startswith(b"%PDF"):
        raise DocumentError("This file does not look like a valid PDF.")
    if extension == ".docx" and not content.startswith(b"PK"):
        raise DocumentError("This file does not look like a valid Word (.docx) document.")
    if extension == ".txt" and b"\x00" in content[:4096]:
        raise DocumentError("This file does not look like a plain text file.")

    return name, extension


def extract_text(content: bytes, extension: str) -> str:
    """Extract plain text from the file. Raises DocumentError on failure."""
    try:
        if extension == ".txt":
            text = content.decode("utf-8", errors="replace")
        elif extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                reader.decrypt("")  # many PDFs are "encrypted" with an empty password
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            import docx

            document = docx.Document(io.BytesIO(content))
            lines = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append(": ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(lines)
        else:
            raise DocumentError(f"Unsupported file type: {extension}")
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError(f"Could not read the document ({type(error).__name__}). The file may be damaged.") from error

    text = text.replace("\r\n", "\n").strip()
    if not text:
        raise DocumentError("No text could be extracted. Scanned image PDFs are not supported (no OCR).")
    return text


# Label patterns for "Label: value" style lines.
FIELD_PATTERNS = {
    "patient_id": r"patient\s*(?:id|number|#)|member\s*(?:id|number|#)",
    "claim_id": r"claim\s*(?:id|number|#)|prescription\s*(?:id|number|#)|rx\s*(?:id|number|#)",
    "medication": r"medication|drug(?:\s*name)?|product",
    "insurance": r"insurance|payer|plan(?:\s*name)?|insurer",
    "rejection_code": r"reject(?:ion)?\s*code|denial\s*code|reason\s*code",
    "rejection_message": r"reject(?:ion)?\s*(?:message|reason|description)|denial\s*reason|message",
    "quantity": r"quantity|qty",
    "date_of_service": r"date\s*of\s*service|service\s*date|fill\s*date|date",
    "prescriber": r"prescriber(?:\s*name)?|prescribing\s*(?:provider|physician)|doctor",
}

DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y", "%B %d, %Y", "%b %d, %Y", "%d %B %Y"]


def _normalize_date(value: str) -> str | None:
    value = value.strip()
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(value, date_format).date().isoformat()
        except ValueError:
            continue
    return None


def extract_fields(text: str) -> dict[str, str | None]:
    """Find common claim fields in the text. Missing fields are returned as None."""
    fields: dict[str, str | None] = {name: None for name in FIELD_PATTERNS}
    for name, label_pattern in FIELD_PATTERNS.items():
        pattern = rf"^\s*(?:{label_pattern})\s*[:#\-]\s*(.+?)\s*$"
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        if match:
            fields[name] = match.group(1)[:500]

    if fields["date_of_service"]:
        fields["date_of_service"] = _normalize_date(fields["date_of_service"])

    # If there is no labelled message, use the first line that mentions a rejection.
    if not fields["rejection_message"]:
        for line in text.splitlines():
            if re.search(r"reject|denied|not covered|authorization|too soon|limit", line, re.IGNORECASE):
                fields["rejection_message"] = line.strip()[:500]
                break

    return fields
