"""
Creates sample_rejection_notice.pdf and sample_rejection_notice.docx for the
upload demo. The generated files are already included in the project; run this
only if you want to re-create them:

    cd backend
    python ../sample_data/generate_sample_documents.py

The PDF is written by hand (a tiny valid PDF), so no extra library is needed.
The DOCX uses python-docx, which is already in backend/requirements.txt.
"""
from pathlib import Path

import docx

OUTPUT_DIR = Path(__file__).resolve().parent

LINES = [
    "SUMMIT SAMPLE HEALTH (FICTIONAL PAYER) - CLAIM REJECTION NOTICE",
    "Sample document for the RxResolveAI student project. All data is fictional.",
    "",
    "Patient ID: PT-20002",
    "Claim ID: RX-70002",
    "Medication: Asthmavent inhaler",
    "Insurance: DemoHealth",
    "Quantity: 2",
    "Date of Service: 2026-09-18",
    "Prescriber: Dr. M. Chen (fictional)",
    "Rejection Code: QL001",
    "Rejection Message: Quantity limit exceeded - maximum quantity 1 inhaler per 30 days",
]


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(path: Path) -> None:
    text_ops = ["BT", "/F1 11 Tf", "14 TL", "50 780 Td"]
    for line in LINES:
        text_ops.append(f"({_pdf_escape(line)}) Tj T*")
    text_ops.append("ET")
    stream = "\n".join(text_ops).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_position = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        output += f"{offset:010d} 00000 n \n".encode()
    output += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_position}\n%%EOF\n".encode()
    path.write_bytes(bytes(output))


def write_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading(LINES[0], level=1)
    for line in LINES[1:]:
        if line:
            document.add_paragraph(line)
    document.save(path)


if __name__ == "__main__":
    write_pdf(OUTPUT_DIR / "sample_rejection_notice.pdf")
    write_docx(OUTPUT_DIR / "sample_rejection_notice.docx")
    print("Created sample_rejection_notice.pdf and sample_rejection_notice.docx in", OUTPUT_DIR)
